
import io
import re
from copy import copy
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="服务站数据自动更新",
    page_icon="📊",
    layout="wide",
)

FIELDS = [
    "25年销货",
    "26年销货单",
    "26年订单",
    "26年合计",
    "期初",
    "已收",
    "总应收",
    "客户余款",
    "总销售",
]

OUTPUT_HEADERS = ["序号", "客户"] + FIELDS


def normalize_name(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r"[\s\u3000]+", "", text)
    return text


def safe_number(value: object) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return 0.0


def find_header_row(ws, max_scan_rows: int = 8):
    for row_idx in range(1, min(ws.max_row, max_scan_rows) + 1):
        values = [
            str(ws.cell(row_idx, col_idx).value).strip()
            if ws.cell(row_idx, col_idx).value is not None else ""
            for col_idx in range(1, ws.max_column + 1)
        ]
        if "客户" in values and "总应收" in values:
            return row_idx, values
    return None, None


def extract_f_records(source_bytes: bytes):
    wb = load_workbook(io.BytesIO(source_bytes), data_only=True)
    records = {}
    duplicate_names = []
    source_errors = []

    for ws in wb.worksheets:
        if ws.title == "汇总":
            continue

        header_row, headers = find_header_row(ws)
        if header_row is None:
            continue

        missing_fields = [field for field in ["客户"] + FIELDS if field not in headers]
        if missing_fields:
            source_errors.append(f"{ws.title} 缺少字段：{', '.join(missing_fields)}")
            continue

        indexes = {header: headers.index(header) + 1 for header in ["客户"] + FIELDS}

        for row_idx in range(header_row + 1, ws.max_row + 1):
            customer = ws.cell(row_idx, indexes["客户"]).value
            if not customer:
                continue

            customer_text = str(customer).strip()
            if not customer_text.upper().startswith("F"):
                continue

            key = normalize_name(customer_text)
            record = {
                "客户": customer_text,
                "来源乡镇": ws.title,
            }
            for field in FIELDS:
                record[field] = safe_number(ws.cell(row_idx, indexes[field]).value)

            if key in records:
                duplicate_names.append(customer_text)
            records[key] = record

    return records, duplicate_names, source_errors


def copy_row_style(ws, source_row: int, target_row: int, max_col: int = 11):
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.fill:
            target.fill = copy(source.fill)
        if source.font:
            target.font = copy(source.font)
        if source.border:
            target.border = copy(source.border)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def find_target_sheet(wb):
    for ws in wb.worksheets:
        for row_idx in range(1, min(ws.max_row, 6) + 1):
            row_values = [ws.cell(row_idx, c).value for c in range(1, min(ws.max_column, 15) + 1)]
            if "客户" in row_values and "26年销货单" in row_values:
                return ws
    return wb.active


def get_existing_order(ws):
    header_row = None
    for row_idx in range(1, min(ws.max_row, 8) + 1):
        values = [ws.cell(row_idx, c).value for c in range(1, min(ws.max_column, 15) + 1)]
        if "客户" in values and "26年销货单" in values:
            header_row = row_idx
            break

    if header_row is None:
        return [], 2

    customer_col = None
    for col in range(1, ws.max_column + 1):
        if ws.cell(header_row, col).value == "客户":
            customer_col = col
            break

    names = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row_idx, customer_col).value
        if not value or str(value).strip() == "合计":
            continue
        names.append(str(value).strip())
    return names, header_row


def build_output(source_bytes: bytes, template_bytes: bytes | None):
    records, duplicates, source_errors = extract_f_records(source_bytes)
    if not records:
        raise ValueError("未在康总数据中找到任何以 F 开头的服务站账户。")

    if template_bytes:
        wb = load_workbook(io.BytesIO(template_bytes))
        ws = find_target_sheet(wb)
        existing_names, header_row = get_existing_order(ws)
    else:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "服务站数据"
        header_row = 2
        existing_names = []

    # 保留原表顺序，新增服务站追加到末尾
    ordered_keys = []
    seen = set()
    for name in existing_names:
        key = normalize_name(name)
        if key in records and key not in seen:
            ordered_keys.append(key)
            seen.add(key)
    for key in records:
        if key not in seen:
            ordered_keys.append(key)
            seen.add(key)

    data_start_row = header_row + 1
    total_row = data_start_row + len(ordered_keys)

    # 清空旧数据区
    for row_idx in range(data_start_row, max(ws.max_row, total_row) + 5):
        for col_idx in range(1, 12):
            ws.cell(row_idx, col_idx).value = None

    # 若模板不存在，建立基础样式
    if not template_bytes:
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=11)
        ws.cell(1, 2).font = Font(size=16, bold=True, color="FFFFFF")
        ws.cell(1, 2).fill = PatternFill("solid", fgColor="1F4E78")
        ws.cell(1, 2).alignment = Alignment(horizontal="center", vertical="center")

        for col_idx, header in enumerate(OUTPUT_HEADERS, start=1):
            cell = ws.cell(header_row, col_idx, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        widths = [8, 30, 14, 14, 14, 14, 12, 14, 14, 14, 14]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

    # 写标题
    ws.cell(1, 2).value = (
        f"销货单统计表（更新至{datetime.now():%Y.%m.%d}，"
        f"{len(ordered_keys)}个服务站）"
    )

    # 写表头，防止模板字段缺失
    for col_idx, header in enumerate(OUTPUT_HEADERS, start=1):
        ws.cell(header_row, col_idx).value = header

    # 复制模板数据行样式
    style_source_row = data_start_row
    if template_bytes and ws.max_row >= style_source_row:
        for row_idx in range(data_start_row, total_row):
            copy_row_style(ws, style_source_row, row_idx)

    # 写入全量服务站数据
    for seq, key in enumerate(ordered_keys, start=1):
        row_idx = data_start_row + seq - 1
        rec = records[key]
        values = [
            seq,
            rec["客户"],
            rec["25年销货"],
            rec["26年销货单"],
            rec["26年订单"],
            rec["26年合计"],
            rec["期初"],
            rec["已收"],
            rec["总应收"],
            rec["客户余款"],
            rec["总销售"],
        ]
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row_idx, col_idx).value = value

    # 合计行样式和公式
    if total_row > data_start_row:
        copy_row_style(ws, max(data_start_row, total_row - 1), total_row)

    ws.cell(total_row, 1).value = None
    ws.cell(total_row, 2).value = "合计"
    for col_idx in range(3, 12):
        letter = get_column_letter(col_idx)
        ws.cell(total_row, col_idx).value = f"=SUM({letter}{data_start_row}:{letter}{total_row - 1})"

    thin = Side(style="thin", color="D6B656")
    for col_idx in range(1, 12):
        cell = ws.cell(total_row, col_idx)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
        cell.font = Font(bold=True, color="7F6000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_idx in range(data_start_row, total_row + 1):
        for col_idx in range(3, 12):
            ws.cell(row_idx, col_idx).number_format = '#,##0.00'

    ws.freeze_panes = f"A{data_start_row}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    preview_rows = []
    for key in ordered_keys:
        rec = records[key]
        preview_rows.append({field: rec[field] for field in ["客户"] + FIELDS})

    return output.getvalue(), pd.DataFrame(preview_rows), duplicates, source_errors


st.title("📊 服务站数据自动更新")
st.caption("上传最新《给康总数据》和上一版服务站表，自动识别全部 F 服务站、全量更新数据并重新计算合计。")

with st.sidebar:
    st.subheader("使用说明")
    st.markdown(
        """
1. 上传最新的《26年给康总数据》  
2. 上传上一版服务站数据（可选，上传后保留原顺序与样式）  
3. 点击“开始更新”  
4. 检查预览后下载新文件
        """
    )

source_file = st.file_uploader(
    "① 上传最新《给康总数据》",
    type=["xlsx"],
    help="系统会自动扫描所有工作表，提取客户名称以 F 开头的服务站。",
)

template_file = st.file_uploader(
    "② 上传上一版服务站数据（可选）",
    type=["xlsx"],
    help="用于保留原客户顺序、列宽和主要样式。",
)

if st.button("🚀 开始更新", type="primary", use_container_width=True):
    if source_file is None:
        st.error("请先上传最新《给康总数据》。")
    else:
        try:
            with st.spinner("正在读取、匹配并更新全部服务站数据……"):
                output_bytes, preview_df, duplicates, source_errors = build_output(
                    source_file.getvalue(),
                    template_file.getvalue() if template_file else None,
                )

            total_count = len(preview_df)
            total_receivable = preview_df["总应收"].sum()
            total_sales = preview_df["总销售"].sum()

            col1, col2, col3 = st.columns(3)
            col1.metric("服务站数量", f"{total_count} 个")
            col2.metric("总应收", f"¥{total_receivable:,.2f}")
            col3.metric("总销售", f"¥{total_sales:,.2f}")

            if duplicates:
                st.warning("发现重复 F 账户，程序已采用最后一次出现的数据：" + "、".join(sorted(set(duplicates))))
            if source_errors:
                st.warning("部分工作表字段不完整：" + "；".join(source_errors))

            st.success("全部服务站数据更新完成。")
            st.dataframe(
                preview_df,
                use_container_width=True,
                hide_index=True,
                column_config={
                    field: st.column_config.NumberColumn(field, format="%.2f")
                    for field in FIELDS
                },
            )

            filename = f"服务站数据_更新至{datetime.now():%Y-%m-%d}.xlsx"
            st.download_button(
                "⬇️ 下载更新后的服务站数据",
                data=output_bytes,
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
            )

        except Exception as exc:
            st.exception(exc)
