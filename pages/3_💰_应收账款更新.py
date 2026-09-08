
import io
import re
from copy import copy
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="应收账款跟进表更新",
    page_icon="💰",
    layout="wide",
)

OUTPUT_HEADERS = [
    "序号",
    "客户名称",
    "销货单",
    "订单",
    "已收款",
    "6.12返点或抹零",
    "总应收",
    "负责人",
    "联系日期",
    "可付款时间",
    "付款状态",
    "下次跟进日期",
]

FOLLOWUP_FIELDS = [
    "负责人",
    "联系日期",
    "可付款时间",
    "付款状态",
    "下次跟进日期",
]

# =========================
# 应收分片区配置
# =========================
AREA_GROUPS = {
    "陈明乡镇": {
        "伏龙泉", "三岗", "永安", "三盛玉", "烧锅", "开安", "刘家",
        "巴吉垒", "龙王", "华家", "前岗", "翁克", "郭家", "黄金",
    },
    "刘春旭": {
        "青山", "哈拉海", "靠山", "万金塔", "高家店", "新农",
        "小城子", "万顺", "黄鱼圈", "杨树林",
    },
    "李艳": {
        "农安镇", "鲍家", "滨河", "柴岗", "德惠", "三宝", "榛柴",
    },
    "小魏": {"合隆"},
    "松原": {"S松原"},
}

AREA_SHEET_ORDER = ["陈明乡镇", "小魏", "李艳", "刘春旭", "松原", "其余"]


def normalize_town_name(value):
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"[\s\u3000]+", "", text)
    return text


def assign_area(source_town):
    """按来源乡镇分配到六个片区；未明确列入的全部归入“其余”。
    因此长春、其他类型、员工以及未来新增未配置片区都会自动进入“其余”。
    """
    town = normalize_town_name(source_town)
    for area_name, towns in AREA_GROUPS.items():
        normalized_towns = {normalize_town_name(x) for x in towns}
        if town in normalized_towns:
            return area_name
    return "其余"

MASTER_FIELDS = [
    "客户",
    "26年销货单",
    "26年订单",
    "已收",
    "6.12返点或抹零",
    "总应收",
    "客户余款",
]


def normalize_name(value):
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r"[\s\u3000]+", "", text)
    return text


def relaxed_name(value):
    """用于兼容 F/S 前缀、活动账户等常见写法。"""
    text = normalize_name(value)
    text = re.sub(r"^[FSfs]+", "", text)
    text = text.replace("参加活动的", "")
    text = text.replace("参加活动", "")
    text = text.replace("活动", "")
    return text


def safe_number(value):
    if value in (None, ""):
        return 0.0
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return 0.0


def find_header_row(ws, required_headers, max_rows=10):
    for row_idx in range(1, min(ws.max_row, max_rows) + 1):
        values = [
            str(ws.cell(row_idx, col_idx).value).strip()
            if ws.cell(row_idx, col_idx).value is not None else ""
            for col_idx in range(1, ws.max_column + 1)
        ]
        if all(header in values for header in required_headers):
            return row_idx, values
    return None, None


def read_master_records(source_bytes):
    """读取《给康总数据》中全部客户。"""
    wb = load_workbook(io.BytesIO(source_bytes), data_only=True)

    records = {}
    duplicate_names = []
    skipped_sheets = []

    for ws in wb.worksheets:
        if ws.title == "汇总":
            continue

        header_row, headers = find_header_row(ws, ["客户", "总应收", "客户余款"])
        if header_row is None:
            skipped_sheets.append(ws.title)
            continue

        index = {
            header: headers.index(header) + 1
            for header in MASTER_FIELDS
        }

        for row_idx in range(header_row + 1, ws.max_row + 1):
            customer = ws.cell(row_idx, index["客户"]).value
            if not customer:
                continue

            customer_text = str(customer).strip()
            if customer_text in {"合计", "总计"}:
                continue

            record = {
                "客户名称": customer_text,
                "销货单": safe_number(ws.cell(row_idx, index["26年销货单"]).value),
                "订单": safe_number(ws.cell(row_idx, index["26年订单"]).value),
                "已收款": safe_number(ws.cell(row_idx, index["已收"]).value),
                "6.12返点或抹零": safe_number(ws.cell(row_idx, index["6.12返点或抹零"]).value),
                "总应收": safe_number(ws.cell(row_idx, index["总应收"]).value),
                "客户余款": safe_number(ws.cell(row_idx, index["客户余款"]).value),
                "来源乡镇": ws.title,
            }

            key = normalize_name(customer_text)
            if key in records:
                duplicate_names.append(customer_text)
            records[key] = record

    return records, duplicate_names, skipped_sheets


def find_receivable_sheet(wb):
    for ws in wb.worksheets:
        header_row, _ = find_header_row(ws, ["客户名称", "总应收"])
        if header_row:
            return ws
    return wb.active


def read_old_followup(template_bytes):
    """读取旧应收跟进表中的人工跟进信息。"""
    wb = load_workbook(io.BytesIO(template_bytes))
    ws = find_receivable_sheet(wb)

    # 原文件带有 Excel Table。更新客户数量后，旧 Table 的 ref / totalsRow
    # 容易与新数据范围不一致，Excel 打开时会提示修复。
    # 本 App 保留单元格样式与筛选数据，但移除底层 Table 定义，确保文件稳定。
    for table_name in list(ws.tables.keys()):
        del ws.tables[table_name]

    header_row, headers = find_header_row(ws, ["客户名称", "总应收"])

    if header_row is None:
        raise ValueError("旧表中未找到“客户名称”和“总应收”表头。")

    col_map = {
        header: headers.index(header) + 1
        for header in headers
        if header
    }

    required_old_headers = ["客户名称", "总应收"] + FOLLOWUP_FIELDS
    missing = [header for header in required_old_headers if header not in col_map]
    if missing:
        raise ValueError("旧应收表缺少必要字段：" + "、".join(missing))

    history = {}
    old_names = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(row_idx, col_map["客户名称"]).value
        if not name or str(name).strip() in {"合计", "总计"}:
            continue
        if str(name).strip().startswith("共") and "户" in str(name):
            continue

        name_text = str(name).strip()
        item = {
            field: ws.cell(row_idx, col_map[field]).value
            for field in FOLLOWUP_FIELDS
        }

        history[normalize_name(name_text)] = item
        history[relaxed_name(name_text)] = item
        old_names.append(name_text)

    return wb, ws, header_row, col_map, history, old_names


def create_blank_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "应收账款跟进表"

    headers = OUTPUT_HEADERS
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(2, col_idx, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [10, 34, 16, 16, 16, 16, 18, 14, 16, 18, 16, 18]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = None
    return wb, ws, 2


def copy_row_style(ws, source_row, target_row, max_col=12):
    if source_row < 1 or source_row > ws.max_row:
        return

    for col_idx in range(1, max_col + 1):
        source = ws.cell(source_row, col_idx)
        target = ws.cell(target_row, col_idx)

        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)

    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def clear_data_area(ws, start_row, end_row, max_col=12):
    for row_idx in range(start_row, end_row + 1):
        for col_idx in range(1, max_col + 1):
            ws.cell(row_idx, col_idx).value = None


def style_total_row(ws, row_idx):
    thin = Side(style="thin", color="D6B656")
    for col_idx in range(1, 13):
        cell = ws.cell(row_idx, col_idx)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
        cell.font = Font(bold=True, color="7F6000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def build_receivable_output(source_bytes, template_bytes=None):
    master_records, duplicates, skipped = read_master_records(source_bytes)

    # 规则：只保留总应收 > 0，且客户余款 <= 0
    eligible = [
        record
        for record in master_records.values()
        if record["总应收"] > 0 and record["客户余款"] <= 0
    ]
    eligible.sort(key=lambda record: record["总应收"], reverse=True)

    if template_bytes:
        wb, ws, header_row, col_map, history, old_names = read_old_followup(template_bytes)
    else:
        wb, ws, header_row = create_blank_template()
        history = {}
        old_names = []
        col_map = {header: idx + 1 for idx, header in enumerate(OUTPUT_HEADERS)}

    # 固定版式：第1行为大标题，第2行为12列表头
    # 如果旧模板识别到的表头不是第2行，后续输出统一按第2行重建。
    header_row = 2

    old_exact_keys = {normalize_name(name) for name in old_names}
    new_customers = [
        record["客户名称"]
        for record in eligible
        if normalize_name(record["客户名称"]) not in old_exact_keys
    ]

    active_names = {normalize_name(record["客户名称"]) for record in eligible}
    removed_customers = [
        name for name in old_names
        if normalize_name(name) not in active_names
    ]

    data_start_row = header_row + 1
    total_row = data_start_row + len(eligible)

    old_max_row = ws.max_row
    clear_data_area(
        ws,
        data_start_row,
        max(old_max_row, total_row + 3),
    )

    # 复制旧表正文行格式
    if template_bytes and old_max_row >= data_start_row:
        for row_idx in range(data_start_row, total_row + 1):
            copy_row_style(ws, data_start_row, row_idx)

    # 强制恢复第1-2行版式：
    # 第1行 = 合并标题；第2行 = 12列表头
    # 先取消任何与第1、2行相交的合并单元格，避免旧模板把第2行吞掉。
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row <= 2 and merged_range.max_row >= 1:
            ws.unmerge_cells(str(merged_range))

    # 旧模板可能把第2行设为隐藏；这里强制显示。
    ws.row_dimensions[1].hidden = False
    ws.row_dimensions[2].hidden = False
    ws.row_dimensions[1].outlineLevel = 0
    ws.row_dimensions[2].outlineLevel = 0
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 36

    # 清空第2行旧内容/旧合并残留，再完整写入12列表头。
    for col_idx in range(1, len(OUTPUT_HEADERS) + 1):
        ws.cell(2, col_idx).value = None

    # 确保表头名称和顺序正确
    for col_idx, header in enumerate(OUTPUT_HEADERS, start=1):
        ws.cell(header_row, col_idx).value = header

    # 统一表头样式：12列全部蓝底、白字、黑体、16号、加粗、居中
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_border_side = Side(style="thin", color="808080")
    header_border = Border(
        left=header_border_side,
        right=header_border_side,
        top=header_border_side,
        bottom=header_border_side,
    )
    for col_idx in range(1, len(OUTPUT_HEADERS) + 1):
        cell = ws.cell(header_row, col_idx)
        cell.fill = header_fill
        cell.font = Font(
            name="SimHei",
            size=16,
            bold=True,
            color="FFFFFF",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = header_border

    new_fill = PatternFill("solid", fgColor="E2F0D9")
    preview_rows = []

    for seq, record in enumerate(eligible, start=1):
        row_idx = data_start_row + seq - 1
        customer = record["客户名称"]

        saved = (
            history.get(normalize_name(customer))
            or history.get(relaxed_name(customer))
            or {}
        )

        values = [
            seq,
            customer,
            record["销货单"],
            record["订单"],
            record["已收款"],
            record["6.12返点或抹零"],
            record["总应收"],
            saved.get("负责人"),
            saved.get("联系日期"),
            saved.get("可付款时间"),
            saved.get("付款状态"),
            saved.get("下次跟进日期"),
        ]

        for col_idx, value in enumerate(values, start=1):
            ws.cell(row_idx, col_idx).value = value

        for money_col in range(3, 8):
            ws.cell(row_idx, money_col).number_format = '#,##0.00'
            ws.cell(row_idx, money_col).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row_idx, 1).alignment = Alignment(horizontal="center", vertical="center")

        # 新客户整行浅绿色标记
        if customer in new_customers:
            for col_idx in range(1, 13):
                ws.cell(row_idx, col_idx).fill = new_fill

        preview_rows.append({
            "序号": seq,
            "客户名称": customer,
            "销货单": record["销货单"],
            "订单": record["订单"],
            "已收款": record["已收款"],
            "6.12返点或抹零": record["6.12返点或抹零"],
            "总应收": record["总应收"],
            "负责人": saved.get("负责人"),
            "联系日期": saved.get("联系日期"),
            "可付款时间": saved.get("可付款时间"),
            "付款状态": saved.get("付款状态"),
            "下次跟进日期": saved.get("下次跟进日期"),
        })

    # 合计行保持用户原表形式：A列“共XX户”，B列“合计”
    ws.cell(total_row, 1).value = f"共{len(eligible)}户"
    ws.cell(total_row, 2).value = "合计"
    for col_idx in range(3, 8):
        letter = get_column_letter(col_idx)
        ws.cell(total_row, col_idx).value = f"=SUM({letter}{data_start_row}:{letter}{total_row - 1})"
        ws.cell(total_row, col_idx).number_format = '#,##0.00'
    style_total_row(ws, total_row)

    # 日期列格式
    for row_idx in range(data_start_row, total_row):
        ws.cell(row_idx, 9).number_format = "yyyy-mm-dd"
        ws.cell(row_idx, 12).number_format = "yyyy-mm-dd"

    ws.freeze_panes = None
    ws.auto_filter.ref = None

    # 顶部标题：合并 A1:L1
    # 上面已经清除了第1-2行冲突的旧合并，这里直接重建标题。
    last_col_letter = get_column_letter(len(OUTPUT_HEADERS))
    ws.merge_cells(f"A1:{last_col_letter}1")

    title_cell = ws["A1"]
    title_cell.value = f"应收账款跟进表 更新至 {datetime.now():%Y.%m.%d}"
    title_cell.fill = PatternFill("solid", fgColor="1F4E78")
    title_cell.font = Font(
        name="SimHei",
        size=20,
        bold=True,
        color="FFFFFF",
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[header_row].height = 30

    # 统一列宽：整体加宽，避免12列表头和正文过于拥挤
    column_widths = {
        "A": 12,   # 序号
        "B": 36,   # 客户名称
        "C": 20,   # 销货单
        "D": 20,   # 订单
        "E": 20,   # 已收款
        "F": 20,   # 6.12返点或抹零
        "G": 22,   # 总应收
        "H": 18,   # 负责人
        "I": 22,   # 联系日期
        "J": 24,   # 可付款时间
        "K": 20,   # 付款状态
        "L": 22,   # 下次跟进日期
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 统一行高：标题、表头、正文、合计行都更舒展
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 34
    for row_idx in range(data_start_row, total_row):
        ws.row_dimensions[row_idx].height = 28
    ws.row_dimensions[total_row].height = 30

    # 最终统一格式：整个应收表完整居中 + 完整细边框
    thin_side = Side(style="thin", color="808080")
    full_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for row in ws.iter_rows(
        min_row=header_row,
        max_row=total_row,
        min_col=1,
        max_col=len(OUTPUT_HEADERS),
    ):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = full_border

    # 标题合并区域外围也加边框
    for col_idx in range(1, len(OUTPUT_HEADERS) + 1):
        ws.cell(1, col_idx).border = full_border

    # 保存前最后一次强制确认第2行12列表头存在且可见
    ws.row_dimensions[2].hidden = False
    ws.row_dimensions[2].height = 36
    for col_idx, header in enumerate(OUTPUT_HEADERS, start=1):
        cell = ws.cell(2, col_idx)
        cell.value = header
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(
            name="SimHei",
            size=16,
            bold=True,
            color="FFFFFF",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = full_border

    # =========================
    # 生成六个分 Sheet
    # =========================
    # 当前应收账款跟进表继续作为总 Sheet，不改变。
    # 每次重新生成六个分表，避免旧数据残留。
    for sheet_name in AREA_SHEET_ORDER:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

    # 将总表中每个符合应收条件的客户，按《给康总数据》的来源乡镇分组。
    area_rows = {sheet_name: [] for sheet_name in AREA_SHEET_ORDER}

    for seq, record in enumerate(eligible, start=1):
        customer = record["客户名称"]
        saved = (
            history.get(normalize_name(customer))
            or history.get(relaxed_name(customer))
            or {}
        )

        area_name = assign_area(record.get("来源乡镇", ""))
        area_rows[area_name].append({
            "客户名称": customer,
            "销货单": record["销货单"],
            "订单": record["订单"],
            "已收款": record["已收款"],
            "6.12返点或抹零": record["6.12返点或抹零"],
            "总应收": record["总应收"],
            "负责人": saved.get("负责人"),
            "联系日期": saved.get("联系日期"),
            "可付款时间": saved.get("可付款时间"),
            "付款状态": saved.get("付款状态"),
            "下次跟进日期": saved.get("下次跟进日期"),
        })

    # 各片区内部继续按总应收从高到低排序。
    for sheet_name in AREA_SHEET_ORDER:
        area_rows[sheet_name].sort(key=lambda item: item["总应收"], reverse=True)

        sub_ws = wb.create_sheet(sheet_name)

        # 标题行 A1:L1
        last_col_letter = get_column_letter(len(OUTPUT_HEADERS))
        sub_ws.merge_cells(f"A1:{last_col_letter}1")
        sub_title = sub_ws["A1"]
        sub_title.value = f"{sheet_name} 应收账款跟进表 更新至 {datetime.now():%Y.%m.%d}"
        sub_title.fill = PatternFill("solid", fgColor="1F4E78")
        sub_title.font = Font(
            name="SimHei",
            size=20,
            bold=True,
            color="FFFFFF",
        )
        sub_title.alignment = Alignment(horizontal="center", vertical="center")
        sub_ws.row_dimensions[1].height = 38

        # 第二行12列表头
        for col_idx, header in enumerate(OUTPUT_HEADERS, start=1):
            cell = sub_ws.cell(2, col_idx, header)
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.font = Font(
                name="SimHei",
                size=16,
                bold=True,
                color="FFFFFF",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = full_border
        sub_ws.row_dimensions[2].height = 36

        # 列宽与总表一致
        sub_column_widths = {
            "A": 12,
            "B": 36,
            "C": 20,
            "D": 20,
            "E": 20,
            "F": 20,
            "G": 22,
            "H": 18,
            "I": 22,
            "J": 24,
            "K": 20,
            "L": 22,
        }
        for col_letter, width in sub_column_widths.items():
            sub_ws.column_dimensions[col_letter].width = width

        sub_data_start = 3

        # 正文
        for sub_seq, item in enumerate(area_rows[sheet_name], start=1):
            r = sub_data_start + sub_seq - 1
            row_values = [
                sub_seq,
                item["客户名称"],
                item["销货单"],
                item["订单"],
                item["已收款"],
                item["6.12返点或抹零"],
                item["总应收"],
                item["负责人"],
                item["联系日期"],
                item["可付款时间"],
                item["付款状态"],
                item["下次跟进日期"],
            ]

            for col_idx, value in enumerate(row_values, start=1):
                cell = sub_ws.cell(r, col_idx, value)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
                cell.border = full_border

            for money_col in range(3, 8):
                sub_ws.cell(r, money_col).number_format = '#,##0.00'

            sub_ws.cell(r, 9).number_format = "yyyy-mm-dd"
            sub_ws.cell(r, 12).number_format = "yyyy-mm-dd"
            sub_ws.row_dimensions[r].height = 28

        # 合计行
        sub_total_row = sub_data_start + len(area_rows[sheet_name])
        sub_ws.cell(sub_total_row, 1).value = f"共{len(area_rows[sheet_name])}户"
        sub_ws.cell(sub_total_row, 2).value = "合计"

        for col_idx in range(3, 8):
            letter = get_column_letter(col_idx)
            if area_rows[sheet_name]:
                sub_ws.cell(sub_total_row, col_idx).value = (
                    f"=SUM({letter}{sub_data_start}:{letter}{sub_total_row - 1})"
                )
            else:
                sub_ws.cell(sub_total_row, col_idx).value = 0
            sub_ws.cell(sub_total_row, col_idx).number_format = '#,##0.00'

        style_total_row(sub_ws, sub_total_row)
        # style_total_row 现在按12列处理。
        sub_ws.row_dimensions[sub_total_row].height = 30

        # 分表不冻结、不自动筛选，保持与总表一致。
        sub_ws.freeze_panes = None
        sub_ws.auto_filter.ref = None

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return {
        "bytes": output.getvalue(),
        "preview": pd.DataFrame(preview_rows),
        "new_customers": new_customers,
        "removed_customers": removed_customers,
        "duplicates": duplicates,
        "skipped_sheets": skipped,
        "customer_count": len(eligible),
        "total_receivable": sum(record["总应收"] for record in eligible),
    }


st.title("💰 应收账款跟进表自动更新")
st.caption(
    "这是独立的应收账款 App，不处理服务站数据。"
    "上传最新《给康总数据》和上一版应收账款跟进表，即可自动生成新表。"
)

with st.sidebar:
    st.subheader("更新规则")
    st.markdown(
        """
- 只保留 **总应收 > 0**
- 只保留 **客户余款 ≤ 0**
- 按总应收从高到低排序
- 自动新增符合条件的新客户
- 自动删除不再符合条件的客户
- 保留人工跟进信息：
  - 负责人
  - 联系日期
  - 可付款时间
  - 付款状态
  - 下次跟进日期
- 新客户整行浅绿色标记
- 自动更新“共XX户”和应收合计
- 不冻结窗口
- 输出文件不保留易损坏的 Excel Table 结构，避免打开时提示修复
- 下载文件中：
  - “应收账款跟进表”为总表
  - 自动生成：陈明乡镇、小魏、李艳、刘春旭、松原、其余 六个分表
  - 未配置到指定片区的乡镇、长春、其他类型、员工自动归入“其余”
        """
    )

source_file = st.file_uploader(
    "① 上传最新《给康总数据》",
    type=["xlsx"],
    help="例如：26年给康总数据_更新至2026-07-27.xlsx",
)

template_file = st.file_uploader(
    "② 上传上一版《应收账款跟进表》",
    type=["xlsx"],
    help="用于保留负责人、联系日期、付款状态等人工填写内容。",
)

if st.button("🚀 开始更新应收账款跟进表", type="primary", use_container_width=True):
    if source_file is None:
        st.error("请先上传最新《给康总数据》。")
    elif template_file is None:
        st.error("请上传上一版《应收账款跟进表》，否则无法保留人工跟进信息。")
    else:
        try:
            with st.spinner("正在读取最新数据、匹配客户并保留跟进信息……"):
                result = build_receivable_output(
                    source_file.getvalue(),
                    template_file.getvalue(),
                )

            c1, c2, c3 = st.columns(3)
            c1.metric("应收客户数", f"{result['customer_count']} 户")
            c2.metric("总应收", f"¥{result['total_receivable']:,.2f}")
            c3.metric("新增客户", f"{len(result['new_customers'])} 户")

            if result["new_customers"]:
                st.success("本次新增：" + "、".join(result["new_customers"]))

            if result["removed_customers"]:
                with st.expander(f"本次移除 {len(result['removed_customers'])} 户"):
                    st.write("、".join(result["removed_customers"]))

            if result["duplicates"]:
                st.warning(
                    "康总数据中发现重复客户名，已采用最后一次出现的数据："
                    + "、".join(sorted(set(result["duplicates"])))
                )

            st.subheader("更新结果预览")
            st.dataframe(
                result["preview"],
                use_container_width=True,
                hide_index=True,
                column_config={
                    "销货单": st.column_config.NumberColumn("销货单", format="%.2f"),
                    "订单": st.column_config.NumberColumn("订单", format="%.2f"),
                    "已收款": st.column_config.NumberColumn("已收款", format="%.2f"),
                    "6.12返点或抹零": st.column_config.NumberColumn("6.12返点或抹零", format="%.2f"),
                    "总应收": st.column_config.NumberColumn("总应收", format="%.2f"),
                    "联系日期": st.column_config.DateColumn("联系日期", format="YYYY-MM-DD"),
                    "下次跟进日期": st.column_config.DateColumn("下次跟进日期", format="YYYY-MM-DD"),
                },
            )

            filename = f"应收账款跟进表_更新至{datetime.now():%Y-%m-%d}.xlsx"
            st.download_button(
                "⬇️ 下载更新后的应收账款跟进表",
                data=result["bytes"],
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
            )

        except Exception as exc:
            st.exception(exc)
